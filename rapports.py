"""Rapport HTML + classeur Excel."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from modele import Post, nombre, pct

BRUN = "7A5C46"


def _dm(ad: dict) -> str:
    return "\n\n".join(f"{e.get('etape', '')}\n{e.get('message', '')}" for e in ad.get("sequence_dm") or [])


def _mise_en_forme(ws, largeurs: dict, gel="B2"):
    from openpyxl.styles import Alignment, Font, PatternFill
    for c in ws[1]:
        ws.column_dimensions[c.column_letter].width = largeurs.get(c.value, 34 if str(c.value).startswith("Hook adapté") else 15)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BRUN)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = gel
    ws.auto_filter.ref = ws.dimensions


def export_excel(posts: list[Post], chemin: Path, cfg: dict, voix_client: dict, decouvertes: list[dict], appr: dict):
    import pandas as pd
    from openpyxl.styles import Font

    planning, analyse = [], []
    for p in posts:
        ad = p.adaptation or {}
        hooks = ad.get("hooks_adaptes") or []
        planning.append({
            "Créneau": p.creneau, "Mécanique": ad.get("mecanique", ""), "Hook retenu": hooks[0] if hooks else "",
            "Texte écran": ad.get("texte_ecran_adapte", ""), "CTA": ad.get("cta_conversation", ""),
            "Tourné": "", "Publié": "", "Vues J+2": "", "Commentaires mot-clé": "", "Conversations DM": "",
            "Clients": "", "Inspiration": p.url,
        })
        analyse.append({
            "Créneau": p.creneau, "Critères OK": "Oui" if p.criteres_ok else "Assoupli", "Mécanique": ad.get("mecanique", ""),
            "Plateforme": p.plateforme, "Compte": "@" + p.compte, "Date": p.date[:10], "Vues": p.vues,
            "Taux d'engagement": pct(p.taux_engagement), "Taux de commentaires": pct(p.taux_commentaires),
            "× médiane compte": p.score_viral, "Lien": p.url,
            "Hook original": ad.get("hook_original", ""), "Texte écran original": ad.get("texte_ecran", ""),
            "Analyse visuelle": ad.get("analyse_visuelle", ""), "Pourquoi ça marche": ad.get("pourquoi_ca_marche", ""),
            "Structure": ad.get("structure", ""), "Ce que dit l'audience": ad.get("ce_que_dit_l_audience", ""),
            **{f"Hook adapté {i + 1}": h for i, h in enumerate(hooks)},
            "Texte écran adapté": ad.get("texte_ecran_adapte", ""), "Script Reel": ad.get("script_reel", ""),
            "CTA conversation": ad.get("cta_conversation", ""), "Séquence DM": _dm(ad),
            "Légende adaptée": ad.get("legende_adaptee", ""), "Format conseillé": ad.get("format_conseille", ""),
            "Idée visuelle": ad.get("idee_visuelle", ""), "Pertinence /5": ad.get("pertinence", ""),
            "Transcription originale": p.transcription,
        })

    o = cfg["objectifs"]
    larg = {"Créneau": 16, "Mécanique": 20, "Hook retenu": 45, "Texte écran": 30, "CTA": 40, "Inspiration": 26,
            "Lien": 26, "Hook original": 40, "Texte écran original": 28, "Analyse visuelle": 45, "Pourquoi ça marche": 50,
            "Structure": 45, "Ce que dit l'audience": 45, "Texte écran adapté": 30, "Script Reel": 70,
            "CTA conversation": 40, "Séquence DM": 70, "Légende adaptée": 70, "Format conseillé": 30,
            "Idée visuelle": 40, "Transcription originale": 70, "Question / douleur": 60, "Hook": 55, "Pourquoi": 60,
            "Compte": 22, "Exemple": 30, "Légende": 50, "Ce qu'on retient": 90}

    with pd.ExcelWriter(chemin, engine="openpyxl") as w:
        pd.DataFrame(planning).to_excel(w, index=False, sheet_name="Planning & suivi")
        pd.DataFrame(analyse).to_excel(w, index=False, sheet_name="Analyse & scripts")

        if voix_client and not voix_client.get("erreur"):
            lignes = [{"Type": t, "Question / douleur": x} for t, cle in
                      (("Question", "questions"), ("Douleur", "douleurs"), ("Objection", "objections"), ("Expression", "expressions"))
                      for x in voix_client.get(cle) or []]
            pd.DataFrame(lignes).to_excel(w, index=False, sheet_name="Voix du client")
            pd.DataFrame([{"Hook": i["hook"], "Pourquoi": i["pourquoi"]} for i in voix_client.get("idees_reels") or []]
                         ).to_excel(w, index=False, sheet_name="Idées de Reels")
        if appr.get("par_mecanique"):
            pd.DataFrame([{
                "Mécanique": l["mecanique"], "Reels analysés": l["n_reels"], "Reels suivis": l["n_suivi"],
                "Vues vs ta moyenne": l.get("vues_vs_moyenne"), "Commentaires vs ta moyenne": l.get("commentaires_vs_moyenne"),
                "Conversations DM moy.": l.get("conversations_moy"), "Indice": l.get("indice"),
            } for l in appr["par_mecanique"]]).to_excel(w, index=False, sheet_name="Ce qui marche chez toi")
        if decouvertes:
            pd.DataFrame([{
                "Plateforme": d["plateforme"], "Compte": "@" + d["compte"], f"Reels ≥ {nombre(cfg['analyse']['vues_min'])} vues": d["reels_au_dessus_seuil"],
                "Vues max": d["vues_max"], "Engagement médian": pct(d["engagement_median"]), "Exemple": d["exemple"], "Légende": d["legende"],
            } for d in decouvertes]).to_excel(w, index=False, sheet_name="Comptes à ajouter")

        for ws in w.sheets.values():
            _mise_en_forme(ws, larg)
        ws = w.sheets["Planning & suivi"]
        n = len(planning) + 1
        ws.cell(n + 2, 1, "TOTAL semaine").font = Font(bold=True)
        for col in (8, 9, 10, 11):  # Vues, mots-clés, conversations, clients
            lettre = ws.cell(1, col).column_letter
            ws.cell(n + 2, col, f"=SUM({lettre}2:{lettre}{n})").font = Font(bold=True)
        ws.cell(n + 3, 1, "Objectif semaine").font = Font(bold=True)
        ws.cell(n + 3, 10, o["conversations_par_jour"] * o["jours_planifies"]).font = Font(bold=True)
        ws.cell(n + 5, 1, "Remplis les colonnes Vues J+2 → Clients puis dépose ce fichier dans le dossier « suivi/ » : "
                          "l'outil apprendra quelles mécaniques te rapportent le plus de conversations.")


def export_html(posts: list[Post], cfg: dict, chemin: Path, stats: dict, voix_client: dict, decouvertes: list[dict], appr: dict,
                media: Path | None = None):
    import base64
    from jinja2 import Environment, FileSystemLoader
    thumbs = {}
    for p in posts:
        if media and p.miniature and (media / p.miniature).exists():
            thumbs[f"{p.plateforme}:{p.id}"] = base64.b64encode((media / p.miniature).read_bytes()).decode()
    env = Environment(loader=FileSystemLoader(str(Path(__file__).parent)), autoescape=True)
    html = env.get_template("gabarit_rapport.html").render(
        posts=posts, niche=cfg["niche"], stats=stats, objectifs=cfg["objectifs"], analyse=cfg["analyse"],
        voix_client=voix_client if voix_client and not voix_client.get("erreur") else {}, decouvertes=decouvertes,
        appr=appr, date=datetime.now().strftime("%d/%m/%Y"), pct=pct, nombre=nombre, thumbs=thumbs,
    )
    chemin.write_text(html, encoding="utf-8")
